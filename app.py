import io
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import streamlit as st

st.set_page_config(
    page_title="Checklist Simulasi Audit Pertamina Way",
    page_icon="📋",
    layout="wide",
)

# --- MASTER CONFIG & CONSTANTS ---
WEIGHT_MAP = {
    "A": 1.00,
    "B": 0.80,
    "C": 0.60,
    "D": 0.40,
    "E": 0.20,
    "F": 0.00,
    "X": 1.00,
}

ALL_CHOICES = ["A", "B", "C", "D", "E", "F", "X"]

CHECKLIST_DATA = {
    "Elemen 1: Skilled Staff & Services (30)": {
        "Sub-Elemen 1.1. Kebersihan dan Penampilan (10)": {
            "1.1.1 Seragam (5)": [
                (
                    "1.1.1.a",
                    "Seluruh operator memakai seragam sesuai standar Pertamina (rancangan serupa, dikancing, baju, celana, sepatu safety warna hitam dan kantong uang minimum 1 tiap pulau pompa)",
                    "A/F",
                    1.25,
                ),
                (
                    "1.1.1.b",
                    "Nama operator dan nomor SPBU tertera dan jelas terbaca",
                    "A/F",
                    1.25,
                ),
                (
                    "1.1.1.c",
                    "Seragam dalam keadaaan bersih dan berkondisi baik",
                    "A-F",
                    1.25,
                ),
                (
                    "1.1.1.d",
                    "Operator tidak membawa telepon genggam (HP) di area pulau pompa selama bertugas",
                    "A/F",
                    1.25,
                    "HP",
                ),
            ],
            "1.1.2 Penampilan & Pemberian Hak (5)": [
                (
                    "1.1.2.a",
                    "Seluruh Operator berpenampilan rapi (operator pria berambut pendek dan disisir rapi, yang berkumis dan berjanggut tipis atau dicukur bersih, untuk wanita; rambut diikat apabila lebih panjang dari bahu, seluruh petugas berkuku pendek)",
                    "A/F",
                    1.50,
                ),
                (
                    "1.1.2.b",
                    "Seluruh operator menerima upah sesuai aturan Upah Minimum dan mendapatkan benefit jaminan kesehatan dan ketenagakerjaan",
                    "A/B/C/F",
                    2.00,
                ),
                (
                    "1.1.2.c",
                    "Pekerja mendapat bagian sesuai haknya dari Program Reward PT Pertamina Patra Niaga",
                    "A/F/X",
                    1.50,
                ),
            ],
        },
        "Sub-Elemen 1.2 Prosedur Pelayanan (20)": {
            "1.2 Prosedur Pelayanan": [
                (
                    "1.2.a",
                    "Pelanggan disambut dengan sopan serta salam (Selamat pagi, siang, sore, malam), ditawarkan produk JBU Top Tier yang tersedia di pulau pompa dan ditanya jenis BBM yang dibutuhkan",
                    "A/C/F",
                    2.00,
                    "SALAM",
                ),
                (
                    "1.2.b",
                    "Operator mengingatkan & memastikan mesin kendaraan konsumen dalam keadaan mati saat pengisian BBM serta memasang stick cone di depan kendaraan konsumen (khusus roda 4)",
                    "A/F",
                    3.00,
                ),
                (
                    "1.2.c",
                    "Setelah pilihan BBM ditetapkan, pelanggan diperlihatkan bahwa penunjuk angka meter dimulai dari angka ’nol’ (X untuk SPBU Self Service)",
                    "A/C/F/X",
                    3.00,
                    "NOL",
                ),
                (
                    "1.2.d",
                    "Pengisian BBM dilakukan secara hati-hati untuk mencegah tumpahnya BBM yang bisa merusak kendaraan (X untuk SPBU Self Service)",
                    "A/F/X",
                    3.00,
                ),
                (
                    "1.2.e",
                    "Operator menawarkan pembayaran menggunakan aplikasi MyPertamina (khusus untuk pelanggan roda empat atau lebih)",
                    "A/F",
                    2.50,
                ),
                (
                    "1.2.f",
                    "Operator mengkonfirmasi harga total dan jumlah uang yang diterima kepada pelanggan",
                    "A-F",
                    2.50,
                ),
                (
                    "1.2.g",
                    "Operator menyerahkan kuitansi/Struk dan memberitahukan jumlah uang kembalian/Self Service : Kasir memberikan uang kembalian sesuai Nota (X untuk SPBU Self Service)",
                    "A/B/C/D/E/F/X",
                    2.00,
                ),
                (
                    "1.2.h",
                    "Operator mengucapkan terima kasih kepada pelanggan atas kunjungannya",
                    "A/C/F",
                    1.00,
                    "TRIMS",
                ),
                (
                    "1.2.i",
                    "Tersedia informasi Call Center Layanan Pelanggan",
                    "A/F",
                    1.00,
                ),
            ]
        },
    },
    "Elemen 2: Exact Quality & Quantity (30)": {
        "Sub-Elemen 2.1: Peralatan (7)": {
            "2.1 Peralatan": [
                (
                    "2.1.a",
                    "Dispenser Unit disegel dan disertifikasi oleh Dinas Metrologi (masa kalibrasi berlaku, segel pada dispenser unit dan sertifikat tersedia)",
                    "A/F",
                    2.50,
                ),
                (
                    "2.1.b",
                    "SPBU memperbaharui secara berkala catatan Totalizer Dispenser Unit BBM yang terdapat di P-Insyst (Jika SPBU belum terdigitalisasi maka diperbolehkan menggunakan catatan manual)",
                    "A/F",
                    2.50,
                ),
                (
                    "2.1.c",
                    "Seluruh peralatan Q&Q tersedia dan dalam kondisi baik",
                    "A/F",
                    2.00,
                ),
            ]
        },
        "Sub-Elemen 2.2 Prosedur Monitoring (23)": {
            "2.2 Monitoring": [
                (
                    "2.2.a",
                    "Tidak ditemukan tanda-tanda manipulasi pada dispenser unit : segel-segel Dispenser, flow meter, dan digital LED yang dapat mempengaruhi ketidakwajaran takaran",
                    "A/F",
                    1.50,
                    "DU DISPENSER",
                ),
                (
                    "2.2.b",
                    "Sampel 2 (dua) pengiriman terakhir dari tiap jenis BBM disimpan dalam kontainer aluminium (1 x 1 liter untuk tiap kompartemen)",
                    "A/F",
                    0.70,
                ),
                (
                    "2.2.c",
                    "Kaleng Sampel disegel dan Label Sampel terisi lengkap",
                    "A/F",
                    0.60,
                ),
                (
                    "2.2.d",
                    "Tersedia Display sampel BBM yang sesuai standar Pertamina",
                    "A/F",
                    0.25,
                ),
                (
                    "2.2.e",
                    "Tidak ditemukan air dalam tangki-tangki timbun",
                    "A/F",
                    1.00,
                ),
                (
                    "2.2.f",
                    "Berat Jenis (densitas) Pertalite (Oktan 90) diukur dengan benar",
                    "A/F/X",
                    0.80,
                ),
                (
                    "2.2.g",
                    "Berat Jenis (densitas) Pertamax (Oktan 92) diukur dengan benar",
                    "A/F/X",
                    0.80,
                ),
                (
                    "2.2.h",
                    "Berat Jenis (densitas) Pertamax Green (Oktan 95) diukur dengan benar",
                    "A/F/X",
                    0.80,
                ),
                (
                    "2.2.i",
                    "Berat Jenis (densitas) Pertamax Turbo (Oktan 98) diukur dengan benar",
                    "A/F/X",
                    0.80,
                ),
                (
                    "2.2.j",
                    "Berat Jenis (densitas) Bio Solar/Solar diukur dengan benar",
                    "A/F/X",
                    0.80,
                ),
                (
                    "2.2.k",
                    "Berat Jenis (densitas) Pertamina Dex diukur dengan benar",
                    "A/F/X",
                    0.85,
                ),
                (
                    "2.2.l",
                    "Berat Jenis (densitas) Dexlite diukur dengan benar",
                    "A/F/X",
                    0.85,
                ),
                (
                    "2.2.m",
                    "Volume BBM yang dikeluarkan dari nozzle akurat",
                    "A/B/C/F",
                    9.25,
                    "UJI PETIK",
                ),
                (
                    "2.2.n",
                    "Catatan stok harian disimpan dan selalu di-update",
                    "A/C/F",
                    1.00,
                ),
                (
                    "2.2.o",
                    "Catatan kualitas harian dan Pemeriksaan Visual tersedia",
                    "A/C/F",
                    1.00,
                ),
                (
                    "2.2.p",
                    "Tanda terima (Surat Pengantar Pengiriman/LO) diarsipkan",
                    "A/F",
                    1.00,
                ),
                (
                    "2.2.q",
                    "Semua produk JBU yang ditawarkan tersedia dan sesuai",
                    "A/F",
                    1.00,
                ),
            ]
        },
    },
    "Elemen 3: Reliable Facilities & Safety (20)": {
        "Sub-Elemen 3.1: Kebersihan harian (14.5)": {
            "3.1.1 Halaman Depan (6.5)": [
                (
                    "3.1.1.a",
                    "Driveway/ Pelataran pengisian BBM dalam kondisi bersih",
                    "A-F",
                    1.00,
                ),
                (
                    "3.1.1.b",
                    "Pulau pompa dan kolom-kolom kanopi dalam kondisi bersih",
                    "A-F",
                    0.75,
                ),
                (
                    "3.1.1.c",
                    "Dispenser Unit BBM dalam kondisi bebas corat-coret",
                    "A-F",
                    0.60,
                ),
                ("3.1.1.d", "SPBU tersedia Oil Spill Kit", "A/F", 0.20),
                (
                    "3.1.1.e",
                    "Totem/Signboard, Lisplang, Kanopi, Rambu bersih tanpa rusak",
                    "A-F",
                    0.50,
                ),
                (
                    "3.1.1.f",
                    "Seluruh lampu penerangan di bawah kanopi berfungsi",
                    "A-F",
                    0.40,
                ),
                (
                    "3.1.1.g",
                    "Seluruh lampu halaman SPBU berfungsi dengan baik",
                    "A-F",
                    0.30,
                ),
                (
                    "3.1.1.h",
                    "Area penyimpanan BBM (area tangki timbun) bersih",
                    "A-F",
                    0.30,
                ),
                (
                    "3.1.1.i",
                    "Pelataran/tempat pembongkaran BBM dalam kondisi bersih",
                    "A-F",
                    0.20,
                ),
                (
                    "3.1.1.j",
                    "Tutup lubang pengisian BBM (Oil Sump) dalam kondisi bersih",
                    "A-F",
                    0.30,
                ),
                (
                    "3.1.1.k",
                    "Lubang pengisian BBM (Oil Sump) bersih dari air/sampah",
                    "A/C/F",
                    0.45,
                ),
                (
                    "3.1.1.l",
                    "Oil Catcher (Jebakan minyak) sesuai standar",
                    "A/F",
                    0.30,
                ),
                ("3.1.1.m", "Taman dalam kondisi bersih", "A/F/X", 0.60),
                (
                    "3.1.1.n",
                    "Lampu penerangan taman berfungsi dengan baik",
                    "A/F/X",
                    0.60,
                ),
            ],
            "3.1.2 Toilet (4.5)": [
                (
                    "3.1.2.a",
                    "Toilet sesuai standar PT Pertamina Patra Niaga",
                    "A/C",
                    1.50,
                ),
                (
                    "3.1.2.b",
                    "Toilet dan wastafel dalam kondisi bersih, terpelihara",
                    "A/C/F",
                    1.20,
                ),
                (
                    "3.1.2.c",
                    "Tidak ada kerusakan yang terlihat pada perangkat toilet",
                    "A/F",
                    0.80,
                ),
                (
                    "3.1.2.d",
                    "Akses mudah menuju toilet disertai tanda penunjuk",
                    "A/F",
                    0.50,
                ),
                (
                    "3.1.2.e",
                    "Toilet berpenampakan cukup dan tersedia perlengkapan",
                    "A/F",
                    0.30,
                ),
                (
                    "3.1.2.f",
                    "Tempat sampah dan keset tersedia pada setiap toilet",
                    "A/F",
                    0.20,
                ),
            ],
            "3.1.3 Fasilitas Ibadah (1)": [
                (
                    "3.1.3.a",
                    "Fasilitas ibadah sesuai standar PT Pertamina",
                    "A/C/X",
                    0.30,
                ),
                (
                    "3.1.3.b",
                    "Area untuk wudhu dan musala dalam keadaan bersih",
                    "A/C/F/X",
                    0.30,
                ),
                (
                    "3.1.3.c",
                    "Tidak ada kerusakan yang terlihat pada musala",
                    "A/C/F/X",
                    0.20,
                ),
                (
                    "3.1.3.d",
                    "Alat perangkat salat tersedia di musala, dan rapi",
                    "A/C/F/X",
                    0.20,
                ),
            ],
            "3.1.4 Aspek HSSE (2.5)": [
                (
                    "3.1.4.a",
                    "Tersedia alat pemadam api ringan (APAR) dengan jumlah sesuai",
                    "A/F",
                    0.15,
                    "APAR",
                ),
                (
                    "3.1.4.b",
                    "Tersedia minimal 2 (dua) buah alat pemadam api roda (APAB)",
                    "A/F",
                    0.15,
                    "APAB",
                ),
                (
                    "3.1.4.c",
                    "Alat pemadam api memiliki kartu inspeksi, isi sesuai",
                    "A/F",
                    0.15,
                    "Masa BERLAKU APAR/APAB",
                ),
                (
                    "3.1.4.d",
                    "Grounding (Kawat penetralan arus listrik) terpasang",
                    "A/F",
                    0.15,
                    "GROUNDING",
                ),
                (
                    "3.1.4.e",
                    "Ducting / saluran pipa dari tangki timbun ke dispenser",
                    "A/F",
                    0.10,
                    "Ducting",
                ),
                (
                    "3.1.4.f",
                    "Seluruh nozzle terpasang Breakaway Valve",
                    "A/F",
                    0.10,
                    "BREAK WAY",
                ),
                (
                    "3.1.4.g",
                    "Dispensing Sump & Impact Valve terpasang dengan baik",
                    "A/F",
                    0.10,
                    "IMPACT VALVE",
                ),
                (
                    "3.1.4.h",
                    "Tidak terdapat kegiatan lain (jualan) di zona bahaya",
                    "A/F",
                    0.10,
                ),
                (
                    "3.1.4.i",
                    "Terdapat stop kontak di kanopi yang sesuai standar",
                    "A/F",
                    0.10,
                ),
                (
                    "3.1.4.j",
                    "Emergency shut down tersedia di setiap pulau pompa",
                    "A/F",
                    0.10,
                    "EMERGENCY SHUT DOWN",
                ),
                (
                    "3.1.4.k",
                    "Junction box kabel kedap (dispenser dan dekat pulau)",
                    "A/F",
                    0.10,
                    "JUNCTION BOX",
                ),
                (
                    "3.1.4.l",
                    "Tidak ada lubang terbuka di area manhole",
                    "A/F",
                    0.10,
                    "ST MANHOLE",
                ),
                (
                    "3.1.4.m",
                    "Tidak terdapat genangan BBM / air di dalam sumuran",
                    "A/F",
                    0.10,
                    "ST AIR",
                ),
                (
                    "3.1.4.n",
                    "Tersedia rambu-rambu / sticker peringatan keselamatan",
                    "A/F",
                    0.10,
                ),
                (
                    "3.1.4.o",
                    "Tersedia daftar nomor telepon penting / emergency",
                    "A/F",
                    0.10,
                ),
                (
                    "3.1.4.p",
                    "Tersedia Surat Ijin Kerja Aman (SIKA) saat perbaikan",
                    "A/F",
                    0.10,
                    "SIKA",
                ),
                (
                    "3.1.4.q",
                    "SPBU tersedia dokumen UKL/UPL yang di-update",
                    "A/C",
                    0.10,
                ),
                ("3.1.4.r", "Instalasi listrik sesuai dengan standar", "A/F", 0.10),
                ("3.1.4.s", "Kotak P3K tersedia di kantor", "A/F", 0.10),
                (
                    "3.1.4.t",
                    "Operator dan pengawas telah terlatih dalam APAR",
                    "A/F",
                    0.10,
                ),
                (
                    "3.1.4.u",
                    "Terdapat minimal 1 (satu) orang petugas SIA/F",
                    "A/F",
                    0.30,
                    "SAFETYMAN",
                ),
            ],
        },
        "Sub-Elemen 3.2: Pemeliharaan berkala atas DU (5.5)": {
            "3.2 Pemeliharaan DU": [
                (
                    "3.2.a",
                    "Catatan pemeliharaan Fasilitas SPBU (house keeping)",
                    "A/F",
                    0.60,
                ),
                (
                    "3.2.b",
                    "Catatan pemeliharaan DU (Dispenser Unit) berkala",
                    "A/F",
                    0.60,
                ),
                (
                    "3.2.c",
                    "Dispenser Unit BBM tidak tampak rusak fisik/bocor",
                    "A/F",
                    0.60,
                ),
                (
                    "3.2.d",
                    "Layar penunjuk (LCD Dispenser) terbaca dengan jelas",
                    "A/F",
                    0.60,
                ),
                (
                    "3.2.e",
                    "Tidak ada kebocoran pada sambungan pipa/dispenser",
                    "A/F",
                    0.60,
                ),
                (
                    "3.2.f",
                    "Semua koneksi listrik di dalam Dispenser Unit aman",
                    "A/F",
                    0.60,
                ),
                (
                    "3.2.g",
                    "Selang pengisian BBM tidak bocor atau terkelupas",
                    "A/F",
                    0.60,
                ),
                (
                    "3.2.h",
                    "Generator terpelihara secara baik dan tidak berisik",
                    "A/F",
                    0.70,
                ),
                (
                    "3.2.i",
                    "Pipa sirkulasi udara tangki timbun (Vent Pipe) baik",
                    "A/F/X",
                    0.60,
                ),
            ]
        },
    },
    "Elemen 4: Visual Format Consistency (10)": {
        "Sub-Elemen 4.1: Identitas Visual Ritel (4)": {
            "4.1 Identitas Visual": [
                (
                    "4.1.a",
                    "Produk Sign (Product Signage) sesuai dengan ketentuan",
                    "A/F",
                    1.00,
                ),
                (
                    "4.1.b",
                    "Totem sesuai dengan standar PERTAMINA",
                    "A/F",
                    1.00,
                ),
                (
                    "4.1.c",
                    "Listplank (Facia) sesuai dengan standar PERTAMINA",
                    "A/F",
                    1.00,
                ),
                (
                    "4.1.d",
                    "Tiang kanopi sesuai dengan standar PERTAMINA",
                    "A/F",
                    1.00,
                ),
            ]
        },
        "Sub-Elemen 4.2: Dispenser Unit (2)": {
            "4.2 Dispenser": [
                (
                    "4.2.a",
                    "Dispenser Unit BBM memiliki warna penutup sesuai",
                    "A/F",
                    1.00,
                ),
                (
                    "4.2.b",
                    "Paduan warna pada Dispenser Unit BBM sesuai",
                    "A/F",
                    1.00,
                ),
            ]
        },
        "Sub-Elemen 4.3: Lain-lain (4)": {
            "4.3 Lain-lain": [
                (
                    "4.3.a",
                    "EDC Digitalisasi dan/atau Tablet MyPertamina tersedia",
                    "A/B/C/F",
                    0.60,
                    "EDC",
                ),
                (
                    "4.3.b",
                    "Petunjuk fasilitas SPBU telah tersedia",
                    "A/F",
                    0.50,
                ),
                (
                    "4.3.c",
                    "Penempatan Signage Tenant sesuai ketentuan",
                    "A/F",
                    0.50,
                ),
                (
                    "4.3.d",
                    "SPBU menggunakan ATG & POS sesuai ketentuan",
                    "A/F",
                    0.60,
                ),
                (
                    "4.3.e",
                    "SPBU dilengkapi CCTV di setiap pulau pompa",
                    "A/F",
                    0.60,
                    "CCTV",
                ),
                (
                    "4.3.f",
                    "Jalur Red Carpet Fast Track (minimal 1 jalur)",
                    "A/C/F",
                    0.40,
                    "RED CARPET",
                ),
                (
                    "4.3.g",
                    "Terdapat Dedicated Operator dengan Rompi Khusus",
                    "A/F",
                    0.40,
                ),
                (
                    "4.3.h",
                    "Posisi jalur Red Carpet Fast Track mudah diakses",
                    "A/F",
                    0.40,
                ),
            ]
        },
    },
    "Elemen 5: Expansive Product Offer (10)": {
        "Sub-Elemen 5.1: Penawaran BBM (2)": {
            "5.1 BBM": [
                ("5.1.a", "Tersedianya produk Pertamax Turbo", "A/F", 0.35),
                ("5.1.b", "Tersedianya produk Pertamax Green", "A/F", 0.25),
                ("5.1.c", "Tersedianya produk Pertamax", "A/F", 0.30),
                ("5.1.d", "Tersedianya produk Pertamina Dex", "A/F", 0.20),
                ("5.1.e", "Tersedianya produk Dexlite", "A/F", 0.20),
                (
                    "5.1.f",
                    "Tersedianya produk JBU minimum 1 jenis yaitu Solar ke atas",
                    "A/C/F",
                    0.40,
                    "JBU",
                ),
                (
                    "5.1.g",
                    "Realisasi penebusan JBU per-2 (dua) bulan terakhir",
                    "A/C/F",
                    0.20,
                ),
                (
                    "5.1.h",
                    "Kesesuaian materi promo yang terpasang di SPBU",
                    "A/F/X",
                    0.10,
                ),
            ]
        },
        "Sub-Elemen 5.2: Penawaran non-BBM (8)": {
            "5.2 Non-BBM": [
                (
                    "5.2.a",
                    "SPBU tersedia NFR Brand Bright yang sesuai",
                    "A/B/F",
                    1.00,
                ),
                (
                    "5.2.b",
                    "SPBU tersedia NFR Internasional",
                    "A/B/F",
                    0.90,
                    "NFR INT",
                ),
                (
                    "5.2.c",
                    "SPBU tersedia fasilitas Energi Baru Terbarukan (EBT)",
                    "A/B/F",
                    1.00,
                ),
                ("5.2.d", "SPBU tersedia NFR Nasional", "A/B/F", 1.00),
                (
                    "5.2.e",
                    "SPBU tersedia NFR Lokal dan berkontrak legal",
                    "A/B/F",
                    0.90,
                    "NFR LKL",
                ),
                (
                    "5.2.f",
                    "Seluruh NFR di SPBU memiliki izin Prinsip/Kontrak",
                    "A/F",
                    0.80,
                    "Izin Prinsip",
                ),
                (
                    "5.2.g",
                    "SPBU Excellent dengan Opsi A: 1. Tersedia Bright Store",
                    "A/B/F",
                    1.00,
                    "NFR",
                ),
                (
                    "5.2.h",
                    "SPBU tersedia 2 (dua) brand pelumas milik Pertamina",
                    "A/B/F",
                    0.90,
                    "PELUMAS",
                ),
                (
                    "5.2.i",
                    "SPBU tersedia fasilitas pengisian air dan angin ban",
                    "A/F",
                    0.50,
                ),
            ]
        },
    },
}

# --- HELPER FUNCTIONS ---
def parse_item(item):
    code, question, valid_options = item[0], item[1], item[2]
    weight = 1.00
    alert_label = ""
    for elem in item[3:]:
        if isinstance(elem, (int, float)):
            weight = float(elem)
        elif isinstance(elem, str):
            alert_label = elem
    return code, question, valid_options, weight, alert_label

def check_is_valid_option(opt, valid_options):
    if not opt:
        return False
    if valid_options == "A-F":
        return opt in ["A", "B", "C", "D", "E", "F"]
    valid_list = [v.strip() for v in valid_options.split("/")]
    return opt in valid_list

def get_allowed_options(valid_options):
    if valid_options == "A-F":
        return ["A", "B", "C", "D", "E", "F"]
    return [v.strip() for v in valid_options.split("/")]

# --- SESSION STATE INITIALIZATION ---
if "answers" not in st.session_state:
    st.session_state.answers = {}
if "notes" not in st.session_state:
    st.session_state.notes = {}
if "media" not in st.session_state:
    st.session_state.media = {}

# --- HEADER & FORM METADATA ---
st.title("📋 Checklist Simulasi Audit Pertamina Way")

with st.container():
    st.subheader("Informasi SPBU")
    c1, c2 = st.columns(2)
    with c1:
        nomor_spbu = st.text_input("NOMOR SPBU", "0000000")
        provinsi = st.text_input("PROVINSI", "")
        kab_kota = st.text_input("KABUPATEN/KOTA", "")
    with c2:
        tipe_spbu = st.selectbox("TIPE SPBU", ["COCO", "DODO", "CODO"])
        tipe_audit = st.selectbox("TIPE AUDIT", ["Initial", "Good", "Excellent"])
        klas_spbu = st.selectbox("KLAS SPBU", ["Basic", "Good", "Excellent"])

    alamat_spbu = st.text_area("ALAMAT SPBU", "")

st.divider()

# --- AUDIT VALIDATION & SCORE CALCULATION ---
has_critical_failure = False
failed_alert_names = []
total_achieved_score = 0.00
total_max_score = 0.00

for elemen_name, sub_elements in CHECKLIST_DATA.items():
    for sub_name, items_dict in sub_elements.items():
        items_to_loop = []
        if isinstance(items_dict, dict):
            for group_name, item_list in items_dict.items():
                items_to_loop.extend(item_list)
        else:
            items_to_loop = items_dict

        for raw_item in items_to_loop:
            code, question, valid_options, weight, alert_label = parse_item(raw_item)
            total_max_score += weight
            
            selected_opt = st.session_state.answers.get(code, None)
            is_valid = check_is_valid_option(selected_opt, valid_options)
            
            if selected_opt and is_valid:
                multiplier = WEIGHT_MAP.get(selected_opt, 0.00)
                total_achieved_score += (multiplier * weight)

            if alert_label and selected_opt == "F" and is_valid:
                has_critical_failure = True
                failed_alert_names.append(f"{code} ({alert_label})")

total_max_score = round(total_max_score, 2)
total_achieved_score = round(total_achieved_score, 2)

# --- STATUS DASHBOARD METRICS ---
score_percentage = (total_achieved_score / total_max_score * 100) if total_max_score > 0 else 0.0

m1, m2, m3 = st.columns(3)
m1.metric("Total Score", f"{total_achieved_score:.2f} / {total_max_score:.2f}")
m2.metric("Pencapaian (%)", f"{score_percentage:.1f}%")

if has_critical_failure:
    m3.error("STATUS: NOT CERTIFIED 🚨")
    alert_list_str = ", ".join(failed_alert_names)
    st.error(
        f"🚨 **STATUS AUDIT: GAGAL (NOT CERTIFIED)** — Temuan "
        f"NOT CERTIFIED pada item ber-Alert: **[{alert_list_str}]**! "
        f"Pelanggaran ini menggagalkan kelulusan audit."
    )
else:
    m3.success("STATUS: CERTIFIED ✅")
    st.success("✅ **STATUS AUDIT: CERTIFIED** — Seluruh item alert memenuhi standar.")

st.divider()

# --- REUSABLE AUDIT ITEM RENDERER ---
def render_audit_item(raw_item):
    code, question, valid_options, weight, alert_label = parse_item(raw_item)
    allowed_opts = get_allowed_options(valid_options)
    
    selected_opt = st.session_state.answers.get(code, None)
    is_valid = check_is_valid_option(selected_opt, valid_options) if selected_opt else False
    multiplier = WEIGHT_MAP.get(selected_opt, 0.00) if (selected_opt and is_valid) else 0.00
    final_score = multiplier * weight

    mandatory_badge = f" 🔥 **[Alert: {alert_label}]**" if alert_label else ""
    
    col_info, col_score = st.columns([4, 1])
    with col_info:
        st.markdown(f"**{code}**: {question}{mandatory_badge}")
        st.caption(f"Valid Opsi: `{valid_options}` | Bobot: `{weight}`")
    with col_score:
        if selected_opt:
            if is_valid:
                st.markdown(f"**Score:** `{final_score:.2f}`")
            else:
                st.markdown("<span style='color:red;'>**Score:** `0.00 (Invalid)`</span>", unsafe_allow_html=True)
        else:
            st.markdown("**Score:** `-`")

    default_idx = allowed_opts.index(selected_opt) if selected_opt in allowed_opts else None
    
    chosen_opt = st.radio(
        f"Pilih Nilai ({code})",
        options=allowed_opts,
        index=default_idx,
        key=f"radio_{code}",
        horizontal=True,
        label_visibility="collapsed"
    )
    st.session_state.answers[code] = chosen_opt

    if alert_label and chosen_opt == "F" and is_valid:
        st.error(f"🚨 ALERT NOT CERTIFIED [{alert_label}] bernilai F (Menggagalkan Kelulusan)!")

    # Notes & Media input
    c_note, c_media = st.columns([2, 2])
    with c_note:
        st.session_state.notes[code] = st.text_input(
            "Catatan Assessor",
            value=st.session_state.notes.get(code, ""),
            key=f"note_{code}",
            placeholder="Tambah catatan jika ada temuan..."
        )
    with c_media:
        uploaded_file = st.file_uploader(
            "📷 Bukti Foto/Video",
            type=["png", "jpg", "jpeg", "mp4", "mov"],
            key=f"media_{code}"
        )
        if uploaded_file is not None:
            st.session_state.media[code] = uploaded_file.name
            st.session_state[f"file_bytes_{code}"] = uploaded_file.getvalue()

    st.markdown("---")

# --- MAIN FORM RENDERER ---
for elemen_name, sub_elements in CHECKLIST_DATA.items():
    with st.expander(f"📁 **{elemen_name}**", expanded=False):
        for sub_name, items_dict in sub_elements.items():
            st.markdown(f"#### _{sub_name}_")
            if isinstance(items_dict, dict):
                for group_name, item_list in items_dict.items():
                    st.markdown(f"**{group_name}**")
                    for raw_item in item_list:
                        render_audit_item(raw_item)
            else:
                for raw_item in items_dict:
                    render_audit_item(raw_item)

# --- EXCEL EXPORT FUNCTION ---
def generate_full_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Laporan Simulasi Audit"
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1F4E78")
    regular_font = Font(name="Calibri", size=11)
    bold_font = Font(name="Calibri", size=11, bold=True)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
    )

    ws['A1'] = "LAPORAN SIMULASI AUDIT PERTAMINA WAY"
    ws['A1'].font = title_font

    overall_status = f"NOT CERTIFIED - Temuan Alert: {', '.join(failed_alert_names)}" if has_critical_failure else "CERTIFIED"

    metadata = [
        ("Nomor SPBU", nomor_spbu),
        ("Provinsi", provinsi),
        ("Kabupaten/Kota", kab_kota),
        ("Alamat SPBU", alamat_spbu),
        ("Tipe SPBU", tipe_spbu),
        ("Tipe Audit", tipe_audit),
        ("Klas SPBU", klas_spbu),
        ("Total Score", f"{total_achieved_score:.2f} / {total_max_score:.2f} ({score_percentage:.1f}%)"),
        ("Status Hasil Audit", overall_status)
    ]

    r = 3
    for k, v in metadata:
        ws.cell(row=r, column=1, value=k).font = bold_font
        ws.cell(row=r, column=2, value=v).font = regular_font
        r += 1

    r += 1
    headers = ["Kode", "Checklist Item", "Nilai", "Bobot", "Score", "Alert", "Catatan Assessor", "Eviden"]
    ws.row_dimensions[r].height = 25
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=r, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = thin_border

    r += 1
    for elemen_name, sub_elements in CHECKLIST_DATA.items():
        for sub_name, items_dict in sub_elements.items():
            items_to_loop = []
            if isinstance(items_dict, dict):
                for group_name, item_list in items_dict.items():
                    items_to_loop.extend(item_list)
            else:
                items_to_loop = items_dict

            for raw_item in items_to_loop:
                code, question, valid_options, weight, alert_label = parse_item(raw_item)
                alert_text = alert_label if alert_label else "-"

                ws.row_dimensions[r].height = 65

                sel_opt = st.session_state.answers.get(code, "-")
                is_valid = check_is_valid_option(sel_opt, valid_options)
                multiplier = WEIGHT_MAP.get(sel_opt, 0.00) if is_valid else 0.00
                final_score = multiplier * weight

                row_data = [
                    (1, code, align_center),
                    (2, question, align_left),
                    (3, sel_opt, align_center),
                    (4, weight, align_center),
                    (5, final_score, align_center),
                    (6, alert_text, align_center),
                    (7, st.session_state.notes.get(code, ""), align_left)
                ]

                for col_idx, val, align in row_data:
                    cell = ws.cell(row=r, column=col_idx, value=val)
                    cell.border = thin_border
                    cell.font = regular_font
                    cell.alignment = align

                cell_media = ws.cell(row=r, column=8)
                cell_media.border = thin_border
                cell_media.alignment = align_center

                if code in st.session_state.media:
                    file_name = st.session_state.media[code]
                    cell_media.value = file_name

                    if f"file_bytes_{code}" in st.session_state and file_name.lower().endswith(('png', 'jpg', 'jpeg')):
                        try:
                            img_io = io.BytesIO(st.session_state[f"file_bytes_{code}"])
                            img = OpenpyxlImage(img_io)
                            img.width = 75
                            img.height = 75
                            ws.add_image(img, f"H{r}")
                        except Exception:
                            pass
                else:
                    cell_media.value = "-"

                r += 1

    col_widths = {'A': 12, 'B': 45, 'C': 10, 'D': 10, 'E': 12, 'F': 18, 'G': 30, 'H': 25}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

# --- DOWNLOAD SECTION ---
st.subheader("📥 Unduh Laporan Excel")
excel_data = generate_full_excel()
st.download_button(
    label="Download Laporan Audit Lengkap (.xlsx)",
    data=excel_data,
    file_name=f"Laporan_Audit_SPBU_{nomor_spbu}.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    use_container_width=True
)
